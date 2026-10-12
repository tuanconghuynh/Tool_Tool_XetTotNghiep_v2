"""Nhận diện cấu trúc theo tên trường và nhóm tiêu đề; không theo vị trí cột."""
import json
import re
import unicodedata
import sys
from pathlib import Path
from collections import Counter
from urllib.parse import urlsplit
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

LABELS = {'academic': 'Học tập', 'records': 'Hồ sơ', 'english': 'Tiếng Anh', 'tuition': 'Học phí'}
ALIASES = {
 'id': ['MaSV', 'MSSV', 'Mã sinh viên', 'Mã học viên'],
 'academic': ['DatTN', 'Tình trạng học tập', 'Kết quả học tập'],
 'records': ['Tình trạng hồ sơ'], 'english': ['Tình trạng chứng chỉ TA', 'Tình trạng tiếng Anh'],
 'tuition': ['Tình trạng học phí'],
 'debt': ['TongHS1', 'Công nợ', 'Nợ học phí'],
 'missing': ['Hồ sơ thiếu'], 'last': ['HoLotSV', 'Họ lót', 'Họ và tên', 'Họ tên'],
 'first': ['TenSV', 'Tên SV', 'Tên'], 'birth': ['NgaySinhC', 'Ngày sinh'],
 'class': ['MaLop', 'Mã lớp'], 'major': ['TenNganh', 'Tên ngành'],
 'credits': ['TCTong', 'Tổng tín chỉ'], 'gpa': ['DTBTN', 'Điểm trung bình'],
 'rank': ['XLTNTC', 'Xếp loại'], 'downgrade': ['HaBacTN'], 'reason': ['LyDoHBTN'],
 'email': ['EMAIL HV', 'Email', 'EmailSV', 'Email học viên'],
 'subjects': ['MÔN CÒN THIẾU', 'Môn thiếu'],
 'verification': ['Xác minh văn bằng', 'XMVB'],
}
PASS = {'academic': ['Đạt', 'Đủ điều kiện'], 'records': ['Đủ hồ sơ'],
        'english': ['Đạt', 'Đủ'], 'tuition': ['Không nợ HP', 'Không nợ học phí', 'Đã hoàn thành học phí']}
FAIL = {'academic': ['Không đạt', 'Chưa đạt'], 'records': ['Thiếu hồ sơ', 'Không có hồ sơ', 'Miễn môn theo bảng điểm CĐ -> chờ KQ xác minh bảng điểm CĐ'],
        'english': ['Không đạt', 'Chưa đạt', 'Thiếu'], 'tuition': ['Nợ HP', 'Nợ học phí', 'Chưa hoàn thành học phí']}
RESULTS = ['ĐỦ ĐIỀU KIỆN', 'KHÔNG ĐỦ ĐIỀU KIỆN', 'CẦN KIỂM TRA']
RETURN_HEADERS = ['HoLotSV','MSSV','NgaySinhC','MaLop','TenNganh','TCTong','DTBTN','XLTNTC','Trạng thái xét tổng','Tình trạng học phí','Hồ sơ thiếu','Tình trạng chứng chỉ TA','HaBacTN','LyDoHBTN','Xác minh văn bằng']
EMAIL_HEADERS = [
 'ĐỢT TN', 'EMAIL HV', 'MÃ SỐ HV', 'HỌ VÀ TÊN', 'NGÀY SINH', 'MÃ LỚP',
 'TÊN NGÀNH', 'TÍN CHỈ TÍCH LŨY', 'ĐIỂM TB HỆ 4', 'XẾP LOẠI TN',
 'TRẠNG THÁI TỐT NGHIỆP', 'TRẠNG THÁI HỌC PHÍ', 'LINK KTRA HỌC PHÍ',
 'MÔN CÒN THIẾU', 'HỒ SƠ THIẾU', 'TRẠNG THÁI CCNN', 'TRẠNG THÁI HẠ BẬC TN',
 'LÝ DO HẠ BẬC', 'XÁC MINH VĂN BẰNG', 'ĐỢT XÉT TN TIẾP THEO',
 'LINK THÔNG BÁO ĐỢT TN TIẾP THEO', 'TIME NHẬN BĂNG TẠM',
 'TIME NHẬN BẰNG CHÍNH THỨC', 'HẠN ĐÓNG HP', 'HẠN BS HỒ SƠ', 'HẠN BS CNCC',
 'HẠN PHẢN HỒI TỔNG', 'STATUS', 'SENT_TIME', 'NOTE',
]
EMAIL_FIELDS = [
 ('ĐỢT TN', 'Đợt xét TN', True, ''),
 ('LINK KTRA HỌC PHÍ', 'Link tra học phí', True, 'https://hocvien.hutech.edu.vn'),
 ('ĐỢT XÉT TN TIẾP THEO', 'Đợt xét TN tiếp theo', True, ''),
 ('LINK THÔNG BÁO ĐỢT TN TIẾP THEO', 'Link thông báo đợt TN tiếp theo', True, ''),
 ('TIME NHẬN BĂNG TẠM', 'Thời gian nhận bằng tạm', True, ''),
 ('TIME NHẬN BẰNG CHÍNH THỨC', 'Thời gian nhận bằng chính thức', True, ''),
 ('HẠN ĐÓNG HP', 'Hạn đóng học phí', False, ''),
 ('HẠN BS HỒ SƠ', 'Hạn bổ sung hồ sơ', False, ''),
 ('HẠN BS CNCC', 'Hạn bổ sung CCNN', False, ''),
 ('HẠN PHẢN HỒI TỔNG', 'Hạn phản hồi tổng', False, ''),
]

def validate_email_metadata(values):
    cleaned = {key: str(values.get(key) or '').strip() for key, _, _, _ in EMAIL_FIELDS}
    missing = [label for key, label, required, _ in EMAIL_FIELDS if required and not cleaned[key]]
    if missing:
        raise ValueError('Vui lòng nhập các mục bắt buộc:\n' + '\n'.join(missing))
    for key in ['LINK KTRA HỌC PHÍ', 'LINK THÔNG BÁO ĐỢT TN TIẾP THEO']:
        try:
            url = urlsplit(cleaned[key])
            valid = url.scheme in {'http', 'https'} and bool(url.hostname) and not any(c.isspace() for c in cleaned[key])
        except ValueError:
            valid = False
        if not valid:
            raise ValueError(f'{key}: vui lòng nhập đường dẫn http:// hoặc https:// hợp lệ.')
    return cleaned

def norm(v):
    t = unicodedata.normalize('NFD', str(v or '').strip().lower()).replace('đ','d')
    return re.sub(r'\s+', ' ', ''.join(c for c in t if unicodedata.category(c) != 'Mn'))

def config(path=None):
    root = Path(sys.executable).parent if getattr(sys, 'frozen', False) else Path(__file__).parent
    p = Path(path) if path else root / 'cau_hinh.json'
    cfg = json.loads(p.read_text(encoding='utf-8')) if p.exists() else {}
    return {k: {**d, **cfg.get(k, {})} for k,d in [('aliases',ALIASES), ('pass',PASS), ('fail',FAIL)]}

def pending(value):
    return bool(re.search(r'\bxmvb\b|\bcho\b|\bdoi\b|\b(?:dang|chua|can) xac minh\b', norm(value)))

def downgrade_fields(value, reason):
    """Chuyển cờ hạ bậc thành lời thông báo; giữ lý do thực tế trong nguồn."""
    text = norm(value)
    if not text or text in {'0', '0.0', 'khong', 'khong ha bac', 'khong ha bac tot nghiep', 'false', 'no'}:
        return None, None
    try:
        has_downgrade = float(str(value).strip()) > 0
    except (ValueError, TypeError):
        has_downgrade = text in {'co', 'x', 'true', 'yes'} or (
            ('ha bac' in text or re.search(r'\bha \d+ bac\b', text))
            and not any(term in text for term in ['khong', 'chua', 'cho'])
        )
    actual_reason = str(reason).strip() if reason is not None else None
    if not has_downgrade:
        return f'Cần kiểm tra hạ bậc: {value}', actual_reason
    return 'Học viên đã được xét hạ bậc tốt nghiệp theo quy định', actual_reason

def missing_documents(value):
    source = str(value or '')
    certificate = r'\bccnn\b|chung chi (?:ngoai ngu|tieng anh)'
    if not re.search(certificate, norm(source)):
        return source
    parts = re.split(r'[;,\n]+|\s+(?:và|/|\+|-)\s+', source, flags=re.IGNORECASE)
    return '; '.join(p.strip() for p in parts if p.strip()
                    and not re.search(certificate, norm(p)))

def verification_details(ws, row, header, cols):
    details = []
    for col in range(1, ws.max_column + 1):
        value = ws.cell(row, col).value
        explicit = col == cols.get('verification')
        if not pending(value) and not (explicit and value and norm(value) not in
                                      {'da xac minh', 'dat', 'hoan thanh', 'khong', '0'}):
            continue
        label = ws.cell(header, col).value or get_column_letter(col)
        if norm(label) == 'tinh trang':
            label = next((LABELS[k] for k in LABELS if cols[k] == col), group(ws, header, col))
        details.append(f'{label}: {value}')
    return '; '.join(details)

def group(ws, row, col):
    for r in range(row-1, 0, -1):
        v = ws.cell(r,col).value
        if not v:
            for area in ws.merged_cells.ranges:
                if area.min_row <= r <= area.max_row and area.min_col <= col <= area.max_col:
                    v = ws.cell(area.min_row,area.min_col).value
                    break
        if v:
            return norm(v)
    return ''

def detect(ws, cfg):
    candidates = []
    for row in range(1, ws.max_row+1):
        headers = {c: norm(ws.cell(row,c).value) for c in range(1,ws.max_column+1)}
        if not any(v in {norm(a) for a in cfg['aliases']['id']} for v in headers.values()):
            continue
        cols = {}
        for key, aliases in cfg['aliases'].items():
            matches = [c for c,v in headers.items() if v and v in {norm(a) for a in aliases}]
            if len(matches)>1:
                raise ValueError(f'{ws.title}, dòng {row}: nhiều cột cùng tên cho {key}; hãy đổi tên rõ ràng.')
            if matches: cols[key] = matches[0]
        for c,h in headers.items():
            if h != 'tinh trang': continue
            g = group(ws,row,c)
            key = next((k for k,terms in [('academic',['hoc tap']),('records',['ho so']),('english',['tieng anh','chung chi ta']),('tuition',['hoc phi'])] if any(t in g for t in terms)), None)
            if key:
                if key in cols and cols[key] != c:
                    raise ValueError(f'{ws.title}: nhiều cột điều kiện {LABELS[key]}; cần đặt tiêu đề riêng.')
                cols[key] = c
        if all(k in cols for k in ['id',*LABELS]): candidates.append((row,cols))
    if len(candidates)>1: raise ValueError(f'{ws.title}: nhiều bảng dữ liệu; cần tách mỗi bảng thành một sheet.')
    return candidates[0] if candidates else None

def classify(value,key,cfg):
    n = norm(value)
    if not n: return 2
    if n in {norm(v) for v in cfg['fail'][key]}: return 1
    if n in {norm(v) for v in cfg['pass'][key]}: return 0
    return 2

def analyze(path, sheet=None, config_path=None):
    cfg = config(config_path)
    wb = load_workbook(path, data_only=True)
    found = []
    for ws in wb:
        if sheet and ws.title != sheet: continue
        detected = detect(ws,cfg)
        if detected: found.append((ws,*detected))
    if len(found) != 1:
        raise ValueError('Không nhận diện duy nhất bảng xét. Cần có mã SV và đủ 4 nhóm học tập, hồ sơ, tiếng Anh, học phí. Dùng --sheet nếu có nhiều sheet xét; đổi tên cột hoặc cập nhật cau_hinh.json nếu tiêu đề thay đổi.')
    ws,header,cols = found[0]
    records=[]
    ids=set()
    for r in range(header+1,ws.max_row+1):
        sid=ws.cell(r,cols['id']).value
        if sid is None or not str(sid).strip():
            if any(ws.cell(r,cols[k]).value is not None for k in LABELS):
                raise ValueError(f'Dòng {r}: có điều kiện nhưng thiếu mã SV; hãy bổ sung mã hoặc bỏ dòng ghi chú khỏi bảng.')
            continue
        if norm(sid) in {norm(a) for a in cfg['aliases']['id']}: continue
        if str(sid) in ids: raise ValueError(f'Mã SV trùng: {sid}, dòng {r}.')
        ids.add(str(sid))
        values={k:ws.cell(r,c).value for k,c in cols.items()}
        codes={k:classify(values[k],k,cfg) for k in LABELS}
        result=1 if 1 in codes.values() else 2 if 2 in codes.values() else 0
        reasons='; '.join(f'{LABELS[k]}: {values[k] if values[k] is not None else "(trống)"}' for k,v in codes.items() if v != 0)
        records.append({'row':r,'values':values,'codes':codes,'result':result,'reasons':reasons,
                        'verification':verification_details(ws,r,header,cols)})
    if not records: raise ValueError('Bảng không có học viên.')
    return ws.title,header,cols,records

def style(ws):
    ws.freeze_panes='A2'
    ws.auto_filter.ref=ws.dimensions
    for c in ws[1]:
        c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='0B5CAD')
        c.alignment=Alignment(wrap_text=True,vertical='center')
    for col in ws.columns:
        ws.column_dimensions[get_column_letter(col[0].column)].width=min(55,max(14,max(len(str(c.value or '')) for c in col)+2))

def process_excel(input_path, output_dir=None, sheet=None, config_path=None, email_metadata=None):
    if email_metadata is not None:
        email_metadata = validate_email_metadata(email_metadata)
    p=Path(input_path)
    if p.suffix.lower() not in {'.xlsx','.xlsm'}: raise ValueError('Chỉ hỗ trợ .xlsx hoặc .xlsm.')
    title,header,cols,rows=analyze(p,sheet,config_path)
    out=Path(output_dir) if output_dir else p.parent
    out.mkdir(parents=True,exist_ok=True)
    wb=load_workbook(p,keep_vba=p.suffix.lower()=='.xlsm')
    # Không chèn/xóa cột nguồn: giữ nguyên công thức, merge, liên kết và sheet phụ.
    base='KQ_TOOL_V2'
    if base in wb: del wb[base]
    detail=wb.create_sheet(base)
    detail.append(['Dòng nguồn','MSSV','Họ tên',*[f'{LABELS[k]} - Trạng thái' for k in LABELS],*[f'{LABELS[k]} - Mã xét' for k in LABELS],'Kết quả xét tốt nghiệp','Lý do / Cảnh báo','Xác minh văn bằng'])
    for rec in rows:
        v=rec['values']; name=' '.join(str(v.get(k) or '').strip() for k in ['last','first']).strip()
        detail.append([rec['row'],v['id'],name,*[v[k] for k in LABELS],*rec['codes'].values(),RESULTS[rec['result']],rec['reasons'],rec['verification']])
    style(detail)
    output_path=out/f'{p.stem}_KET_QUA{p.suffix}'
    if output_path.resolve()==p.resolve(): raise ValueError('Không được ghi đè file nguồn.')
    wb.save(output_path)
    report=Workbook(); summary=report.active; summary.title='TONG_QUAN'
    counts=Counter(x['result'] for x in rows)
    summary.append(['Chỉ tiêu','Giá trị']); summary.append(['Sheet nguồn',title]);summary.append(['Dòng tiêu đề',header]);summary.append(['Tổng học viên',len(rows)])
    for i,t in enumerate(RESULTS): summary.append([t,counts[i]])
    for k in LABELS: summary.append([f'{LABELS[k]} - Cột nguồn',get_column_letter(cols[k])])
    style(summary)
    d=report.create_sheet('CHI_TIET')
    for row in detail.values: d.append(row)
    style(d)
    warn=report.create_sheet('CANH_BAO'); warn.append(['Dòng','MSSV','Nhóm','Giá trị chưa nhận diện'])
    for rec in rows:
        for k,c in rec['codes'].items():
            if c==2: warn.append([rec['row'],rec['values']['id'],LABELS[k],rec['values'][k] or '(trống)'])
    style(warn)
    report_path=out/f'{p.stem}_BAO_CAO.xlsx'; report.save(report_path)
    ret=Workbook(); data=ret.active;data.title='DATA';data.append(RETURN_HEADERS)
    for rec in rows:
        v=rec['values']; codes=rec['codes']
        data.append([' '.join(str(v.get(k) or '').strip() for k in ['last','first']).strip(),v['id'],*[v.get(k) for k in ['birth','class','major','credits','gpa','rank']],['Du dieu kien','Khong du dieu kien','Can kiem tra'][rec['result']],['Không','Có','Cần kiểm tra'][codes['tuition']],missing_documents(v.get('missing')),['Đủ','Thiếu','Cần kiểm tra'][codes['english']],v.get('downgrade') or None,v.get('reason'),rec['verification']])
    style(data)
    note=ret.create_sheet('Note');note.append(['Quy ước','Nội dung']);note.append(['0 / 1 / 2','Đạt / Không đạt / Chưa nhận diện hoặc trống']);note.append(['Cần kiểm tra','Không đưa vào danh sách đủ điều kiện; xác minh trước khi nhập hệ thống.']);note.append(['Học phí','Theo tình trạng học phí, không suy từ số dư TongHS1.']);style(note)
    return_path=out/f'{p.stem}_TRA_KET_QUA.xlsx';ret.save(return_path)
    email_path = create_email_result_file(p, title, header, rows, out, email_metadata)
    return {'output_path':output_path,'report_path':report_path,'return_result_path':return_path,'email_result_path':email_path,'sheet':title,'header_row':header,'total_count':len(rows),'eligible_count':counts[0],'not_eligible_count':counts[1],'review_count':counts[2]}

def create_email_result_file(path, title, header, rows, output_dir, email_metadata=None):
    source = load_workbook(path, data_only=True)
    ws = source[title]
    headers = {norm(ws.cell(header,c).value): c for c in range(1,ws.max_column+1)
               if ws.cell(header,c).value}
    wb = Workbook(); data = wb.active; data.title = 'DATA'; data.append(EMAIL_HEADERS)
    for rec in rows:
        v = rec['values']; codes = rec['codes']
        values = [ws.cell(rec['row'],headers[norm(h)]).value if norm(h) in headers else None
                  for h in EMAIL_HEADERS[:27]] + [None, None, None]
        values[1] = v.get('email')
        values[2] = v['id']
        values[3] = ' '.join(str(v.get(k) or '').strip() for k in ['last','first']).strip()
        for index,key in [(4,'birth'),(5,'class'),(6,'major'),(7,'credits'),(8,'gpa'),(9,'rank'),
                          (13,'subjects'),(16,'downgrade'),(17,'reason')]:
            values[index] = v.get(key)
        values[10] = ['Du dieu kien','Khong du dieu kien','Can kiem tra'][rec['result']]
        values[11] = ['Đã hoàn thành','Chưa hoàn thành học phí','Cần kiểm tra'][codes['tuition']]
        values[14] = missing_documents(v.get('missing'))
        values[15] = ['Đủ','Thiếu','Cần kiểm tra'][codes['english']]
        values[18] = rec['verification']
        values[16], values[17] = downgrade_fields(v.get('downgrade'), v.get('reason'))
        if email_metadata is not None:
            for key, value in email_metadata.items():
                values[EMAIL_HEADERS.index(key)] = value or None
        elif not values[12]:
            values[12] = 'https://hocvien.hutech.edu.vn'
        data.append(values)
    source.close()
    style(data)
    for row in data.iter_rows(min_row=2):
        for cell in row: cell.alignment = Alignment(wrap_text=True, vertical='center')
    data.row_dimensions[1].height = 48
    note = wb.create_sheet('Note')
    note.append(['Quy ước', 'Nội dung'])
    note.append(['XÁC MINH VĂN BẰNG', 'XMVB/đang chờ giữ nguyên nội dung kèm tên cột nguồn để kiểm tra.'])
    note.append(['HỒ SƠ THIẾU', 'Lấy toàn bộ nội dung Hồ sơ thiếu từ file nguồn, chỉ bỏ phần CCNN. Các mục XMVB/đang chờ khác vẫn giữ; XMVB CCNN xem tại XÁC MINH VĂN BẰNG.'])
    note.append(['TRẠNG THÁI HẠ BẬC TN', 'Nếu có: Học viên đã được xét hạ bậc tốt nghiệp theo quy định. Không hạ bậc thì để trống.'])
    note.append(['LÝ DO HẠ BẬC', 'Lấy lý do thực tế từ LyDoHBTN nếu có hạ bậc. Ví dụ: % học phần thi lại vượt mức cho phép - Học lại 9 TC.'])
    note.append(['Trường trống', 'Bổ sung email, đợt xét, đường dẫn và thời hạn còn thiếu trước khi gửi.'])
    note.append(['Cần kiểm tra', 'Xác minh các dòng Can kiem tra trước khi gửi email.'])
    note.append(['STATUS / SENT_TIME / NOTE', 'Để trống cho tool email quản lý lần gửi mới.'])
    style(note)
    result = Path(output_dir) / f'{path.stem}_TOOL_EMAIL.xlsx'
    wb.save(result); wb.close()
    return result
