import unittest
import tempfile
from pathlib import Path
from openpyxl import Workbook, load_workbook
from engine import analyze, classify, config, process_excel, EMAIL_HEADERS, missing_documents, EMAIL_FIELDS, validate_email_metadata, downgrade_fields

class EngineTests(unittest.TestCase):
    def test_missing_documents_preserves_all_except_ccnn(self):
        original = 'CCCD; Đang chờ XMVB bằng THPT; Học bạ'
        self.assertEqual(missing_documents(original), original)
        self.assertEqual(missing_documents(original + '; CCNN'), original)
        self.assertEqual(missing_documents('CCNN / CCCD / Đang chờ bảng điểm'), 'CCCD; Đang chờ bảng điểm')
        self.assertEqual(missing_documents('XMVB CCNN; Đang chờ XMVB bằng THPT'), 'Đang chờ XMVB bằng THPT')
        with tempfile.TemporaryDirectory() as d:
            p = self.fixture(d); w = load_workbook(p); s = w['Danh sach']
            s.cell(2, 6, 'Hồ sơ thiếu'); s.cell(3, 6, original + '; XMVB CCNN')
            w.save(p); w.close()
            info = process_excel(p)
            for key, cell in [('return_result_path', 'K2'), ('email_result_path', 'O2')]:
                output = load_workbook(info[key])
                self.assertEqual(output['DATA'][cell].value, original)
                output.close()
            email = load_workbook(info['email_result_path'])
            self.assertIn('XMVB CCNN', email['DATA']['S2'].value)
            email.close()

    def test_downgrade_status_and_actual_reason(self):
        reason = '% học phần thi lại vượt mức cho phép - Học lại 9 TC'
        message = 'Học viên đã được xét hạ bậc tốt nghiệp theo quy định'
        for value in [1, 1.0, '1', 'Có', 'X', True]:
            self.assertEqual(downgrade_fields(value, reason), (message, reason))
        for value in [None, '', 0, 0.0, '0', 'Không', False]:
            self.assertEqual(downgrade_fields(value, reason), (None, None))
        self.assertEqual(downgrade_fields(1, None), (message, None))
        self.assertIn('Cần kiểm tra', downgrade_fields('Chờ xét hạ bậc', None)[0])
        with tempfile.TemporaryDirectory() as d:
            p = self.fixture(d); w = load_workbook(p); s = w['Danh sach']
            s.cell(2, 6, 'HaBacTN'); s.cell(2, 7, 'LyDoHBTN')
            s.cell(3, 6, 1); s.cell(3, 7, reason)
            w.save(p); w.close()
            info = process_excel(p)
            email = load_workbook(info['email_result_path'])
            self.assertEqual(email['DATA']['Q2'].value, message)
            self.assertEqual(email['DATA']['R2'].value, reason)
            email.close()

    def email_metadata(self):
        return {key: default or ('10/2026' if required else '')
                for key, _, required, default in EMAIL_FIELDS} | {
                    'LINK THÔNG BÁO ĐỢT TN TIẾP THEO': 'https://www.hutech.edu.vn/thong-bao',
                    'TIME NHẬN BĂNG TẠM': 'Từ ngày 20/10/2026',
                    'TIME NHẬN BẰNG CHÍNH THỨC': 'Theo thông báo của trường',
                }

    def test_required_metadata_and_optional_deadlines(self):
        values = self.email_metadata()
        self.assertEqual(validate_email_metadata(values)['HẠN PHẢN HỒI TỔNG'], '')
        for key, _, required, _ in EMAIL_FIELDS:
            if required:
                with self.subTest(key=key):
                    with self.assertRaises(ValueError):
                        validate_email_metadata(values | {key: '  '})
        with self.assertRaises(ValueError):
            validate_email_metadata(values | {'LINK KTRA HỌC PHÍ': 'khong-phai-link'})

    def test_metadata_applies_to_every_student(self):
        with tempfile.TemporaryDirectory() as d:
            p = self.fixture(d)
            w = load_workbook(p); s = w['Danh sach']
            s.append(['002', 'Đạt', 'Đủ hồ sơ', 'Đạt', 'Không nợ HP'])
            s.cell(2, 6, 'HẠN ĐÓNG HP'); s.cell(3, 6, 'Hạn cũ trong nguồn')
            w.save(p); w.close()
            values = self.email_metadata() | {'HẠN BS HỒ SƠ': '25/10/2026'}
            info = process_excel(p, email_metadata=values)
            email = load_workbook(info['email_result_path'])
            self.assertEqual(email['DATA'].max_row, 3)
            for row in [2, 3]:
                for key, value in values.items():
                    self.assertEqual(email['DATA'].cell(row, EMAIL_HEADERS.index(key)+1).value, value or None)
            email.close()

    def test_invalid_metadata_creates_no_output(self):
        with tempfile.TemporaryDirectory() as d:
            p = self.fixture(d); out = Path(d)/'out'
            with self.assertRaises(ValueError):
                process_excel(p, out, email_metadata={})
            self.assertFalse(out.exists())

    def fixture(self, folder, shift=0, reverse=False):
        w=Workbook();w.active.title='Ghi chu';w.active.append(['Không phải bảng xét'])
        s=w.create_sheet('Danh sach')
        columns=[('MaSV',''),('DatTN','Kết quả học tập'),('Tình trạng','Hồ sơ'),('Tình trạng','Chứng chỉ Tiếng Anh'),('Tình trạng','Học phí')]
        values=['001','Đạt','Đủ hồ sơ','Đạt','Không nợ HP']
        order=list(range(5))[::-1] if reverse else list(range(5))
        for c,i in enumerate(order,1):
            s.cell(1+shift,c,columns[i][1]);s.cell(2+shift,c,columns[i][0]);s.cell(3+shift,c,values[i])
        # Dòng trống có định dạng không được tính thành học viên.
        s.cell(8+shift,1).number_format='@'
        p=Path(folder)/'source.xlsx';w.save(p);return p

    def test_reorder_shift_and_sheet(self):
        with tempfile.TemporaryDirectory() as d:
            for shift,rev in [(0,False),(30,True)]:
                p=self.fixture(d,shift,rev)
                title,header,cols,rows=analyze(p)
                self.assertEqual(title,'Danh sach');self.assertEqual(header,2+shift)
                self.assertEqual(len(rows),1);self.assertEqual(rows[0]['result'],0)
                info=process_excel(p,Path(d)/'out')
                self.assertEqual(info['total_count'],1)
                w=load_workbook(info['output_path'])
                self.assertEqual(w['Danh sach'].cell(3+shift,cols['id']).value,'001')
                self.assertEqual(w['Ghi chu']['A1'].value,'Không phải bảng xét')

    def test_unknown_and_negation(self):
        cfg=config()
        self.assertEqual(classify('Không đạt','academic',cfg),1)
        self.assertEqual(classify('Không nợ HP','tuition',cfg),0)
        self.assertEqual(classify('Nợ HP','tuition',cfg),1)
        self.assertEqual(classify(None,'english',cfg),2)
        self.assertEqual(classify('Chờ xét đạt','academic',cfg),2)

    def test_missing_condition_blocks(self):
        with tempfile.TemporaryDirectory() as d:
            p=self.fixture(d);w=load_workbook(p);w['Danh sach']['E1']='Không xác định';w.save(p)
            with self.assertRaises(ValueError): analyze(p)

    def test_merged_group(self):
        with tempfile.TemporaryDirectory() as d:
            p=self.fixture(d);w=load_workbook(p);s=w['Danh sach']
            s.insert_cols(3);s['C1']='Hồ sơ';s['D1']=None;s.merge_cells('C1:D1');w.save(p)
            self.assertEqual(analyze(p)[3][0]['result'],0)

    def test_email_export_and_verification(self):
        with tempfile.TemporaryDirectory() as d:
            p=self.fixture(d,30,True)
            w=load_workbook(p);s=w['Danh sach']
            for col,header,value in [(6,'Hồ sơ thiếu','CCCD; CCNN; XMVB CCNN'),
                                      (7,'Email','test@example.test'),
                                      (8,'Ghi chú','Đang chờ xác minh bằng THPT')]:
                s.cell(32,col,header);s.cell(33,col,value)
            w.save(p);w.close()
            info=process_excel(p,Path(d)/'out')
            self.assertEqual(info['eligible_count'],1)
            email=load_workbook(info['email_result_path'])
            self.assertEqual([c.value for c in email['DATA'][1]],EMAIL_HEADERS)
            self.assertEqual(email['DATA']['B2'].value,'test@example.test')
            self.assertEqual(email['DATA']['C2'].value,'001')
            self.assertEqual(email['DATA']['O2'].value,'CCCD')
            self.assertIn('XMVB CCNN',email['DATA']['S2'].value)
            self.assertIn('Đang chờ xác minh bằng THPT',email['DATA']['S2'].value)
            self.assertEqual(email['DATA']['L2'].value,'Đã hoàn thành')
            self.assertTrue(all(email['DATA'].cell(2,col).value is None for col in [28,29,30]))
            self.assertEqual(email['DATA'].max_row,2);email.close()
            ret=load_workbook(info['return_result_path'])
            self.assertEqual(ret['DATA']['K2'].value,'CCCD')
            self.assertIn('XMVB CCNN',ret['DATA']['O2'].value);ret.close()
            result=load_workbook(info['output_path'])
            self.assertEqual(result['KQ_TOOL_V2']['N1'].value,'Xác minh văn bằng')
            self.assertIn('XMVB CCNN',result['KQ_TOOL_V2']['N2'].value);result.close()
            self.assertEqual(missing_documents('CCCD và CCNN'),'CCCD')

    def test_email_unknown_is_not_eligible(self):
        with tempfile.TemporaryDirectory() as d:
            p=self.fixture(d);w=load_workbook(p);w['Danh sach']['D3']='Đang XMVB CCNN';w.save(p);w.close()
            info=process_excel(p)
            self.assertEqual(info['review_count'],1)
            email=load_workbook(info['email_result_path'])
            self.assertEqual(email['DATA']['K2'].value,'Can kiem tra')
            self.assertEqual(email['DATA']['P2'].value,'Cần kiểm tra')
            self.assertIn('XMVB CCNN',email['DATA']['S2'].value);email.close()

if __name__=='__main__': unittest.main()
